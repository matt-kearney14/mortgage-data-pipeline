"""Shared pieces of the pipeline: input checks, the event loader, sessions, and
the page-template lookup.

Every stage imports from here, so each of these is defined exactly once:
  - which input files, sheets and columns the pipeline needs (check_inputs)
  - how the event log is read, sorted and cleaned of pre-pilot test users
  - how a path maps to its page template in docs/page_template_coding.csv
  - how events are cut into sessions and how an audio clip finds its page

Nothing in this file names a user, a loan or a count from the current data.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
DOCS = ROOT / "docs"
OUT = ROOT / "output"

EVENTS_FILE = DATA / "talkument_userinteractions.xlsx"
ACCOUNT_FILE = DATA / "talkument_useraccount.xlsx"
APPLICANT_FILE = DATA / "talkument_loan_applicants.xlsx"
BUCKET_FILE = DATA / "talkument_pilot_buckets.xlsx"
LOAN_FILE = DATA / "loan_application_data_partial.csv"
CODING_FILE = DOCS / "Clickstream_path_frequencies_and_coding_scheme.xlsx"
TEMPLATE_FILE = DOCS / "page_template_coding.csv"
DICTIONARY_FILE = OUT / "path_dictionary_extended.csv"

# file -> (sheet or None for csv, required columns)
REQUIRED_INPUTS = {
    EVENTS_FILE: ("user_usage", ["user_hash", "eventdate", "path"]),
    ACCOUNT_FILE: ("users", ["user_hash", "provided_language", "expertise_level",
                             "account_enabled"]),
    APPLICANT_FILE: ("loan_applicants", ["loannumber", "user_hash", "language_preference",
                                         "state", "applicant_email_hash"]),
    BUCKET_FILE: ("pilot_record", ["loan_number", "bucket"]),
    LOAN_FILE: (None, ["Loan_Number", "Loan_Status", "Application_Date", "LE_TIL_Sent_Date",
                       "Lock_Date", "Current_Status_Date", "bucket", "Activated_Talkument",
                       "Borrower_Language_Preference", "Coapplicant", "Credit_Score_Decision",
                       "Interest_Rate", "APR", "HMDA_Loan_Type", "HMDA_Loan_Purpose"]),
    CODING_FILE: ("beta_coding", ["CODING SCHEME"]),
    TEMPLATE_FILE: (None, ["template", "page_group", "content_kind"]),
}

# The professor's flags, in his column order, plus the two document flags the
# template table adds. Audio and AudioMp3 are computed, not looked up.
FLAGS = ["Personalized", "GeneralFinancial", "MortgageRelated", "ProcessRelated",
         "BorrowerMortgageProcessRelated", "LenderMortgageProcessRelated", "LoanTermsRelated",
         "LoanEstimateRelated", "CDRelated", "LEDocument", "CDDocument", "Download", "Video",
         "Goal_to_inform", "Goal_to_Advise"]

USER, TIME, PATH = "user_hash", "eventdate", "path"
DOWNLOAD_PREFIX = "/Download/LoanDocument/"
TRANSLATION_PREFIX = "/translations/"
SESSION_TIMEOUT_MIN = 30      # DEC-N; our assumption, a CLI flag in phase 2

# DEC-AA. Activity before the pilot opened is testing, not borrowers. None means
# "the earliest Application_Date in the loan extract", so new data sets its own.
PILOT_START: pd.Timestamp | None = None


# ============================================================ input checks
def check_inputs(files=None) -> None:
    """Stop with a plain-language message if an input file, sheet or column is
    missing. Reads only header rows, so it costs a second, not a full load."""
    from openpyxl import load_workbook
    problems = []
    for f in (files or REQUIRED_INPUTS):
        sheet, cols = REQUIRED_INPUTS[f]
        rel = f.relative_to(ROOT)
        if not f.exists():
            problems.append(f"- {rel}: file not found")
            continue
        if sheet is None:
            have = list(pd.read_csv(f, nrows=0).columns)
        else:
            wb = load_workbook(f, read_only=True)
            if sheet not in wb.sheetnames:
                problems.append(f"- {rel}: no sheet named '{sheet}' "
                                f"(sheets found: {', '.join(wb.sheetnames)})")
                wb.close()
                continue
            have = [c.value for c in next(wb[sheet].iter_rows(max_row=1))]
            wb.close()
        missing = [c for c in cols if c not in have]
        if missing:
            problems.append(f"- {rel} (sheet '{sheet}'): missing column(s) {', '.join(missing)}"
                            if sheet else f"- {rel}: missing column(s) {', '.join(missing)}")
    if problems:
        sys.exit("The input files do not match what the pipeline expects:\n"
                 + "\n".join(problems)
                 + "\nFile names, sheet names and column names must match the original "
                   "data exactly (see README, 'Running it').")


# ============================================================ events
def pilot_start() -> pd.Timestamp:
    """DEC-AA: the first day of the pilot, from the data unless overridden."""
    if PILOT_START is not None:
        return pd.Timestamp(PILOT_START)
    d = pd.to_datetime(pd.read_csv(LOAN_FILE, usecols=["Application_Date"]).Application_Date,
                       format="%m/%d/%Y", errors="coerce")
    return d.min().normalize()


def load_events() -> tuple[pd.DataFrame, dict]:
    """The event log in spec order (user, time, file row), test users removed.

    Returns (events, info). info records what was removed so each stage can
    report it rather than re-derive it.
    """
    ev = pd.read_excel(EVENTS_FILE, sheet_name="user_usage", dtype={PATH: str, USER: str})
    try:
        ev[TIME] = pd.to_datetime(ev[TIME], format="%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError) as e:
        sys.exit(f"{EVENTS_FILE.name}: eventdate must be 'YYYY-MM-DD HH:MM:SS' ({e}).")
    ev[PATH] = ev[PATH].astype(str)
    # CLAUDE.md hygiene 3: the file row is the tiebreaker, captured before sorting
    ev["_source_row"] = np.arange(len(ev))
    ev = ev.sort_values([USER, TIME, "_source_row"], kind="mergesort").reset_index(drop=True)

    # DEC-AA: a user whose FIRST event predates the pilot is a tester. Whole users
    # are dropped, not just their early rows, so no tester leaves a partial trail.
    start = pilot_start()
    first = ev.groupby(USER)[TIME].transform("min")
    test = first < start
    info = dict(pilot_start=start, raw_events=len(ev), raw_users=int(ev[USER].nunique()),
                test_users=sorted(ev.loc[test, USER].unique()),
                test_events=int(test.sum()))
    return ev[~test].reset_index(drop=True), info


# ============================================================ templates
def _load_templates() -> pd.DataFrame:
    t = pd.read_csv(TEMPLATE_FILE, dtype=str, keep_default_na=False)
    bad = t.template.duplicated()
    if bad.any():
        sys.exit(f"{TEMPLATE_FILE.name}: duplicate template(s) {list(t.template[bad])}")
    return t


TEMPLATES = _load_templates() if TEMPLATE_FILE.exists() else pd.DataFrame(
    columns=["template", "page_group", "content_kind"])


def fold_path(path: str) -> str:
    """A path's template key: query-free, no trailing slash, and the numbered
    variants the app serves (-1, -2) folded onto their page."""
    p = path.split("?")[0]
    if len(p) > 1:
        p = p.rstrip("/")
    return re.sub(r"-\d+$", "", p)


def template_of(path: str) -> str | None:
    """The template row a path belongs to, or None. Exact keys win over prefix
    keys (ending in *); among prefixes the longest wins."""
    if path.endswith(".mp3"):
        return None
    key = fold_path(path)
    exact = set(TEMPLATES.template)
    if key in exact:
        return key
    best = None
    for t in TEMPLATES.template:
        if t.endswith("*") and path.startswith(t[:-1]) and (best is None or len(t) > len(best)):
            best = t
    return best


# ============================================================ sessions
def session_ids(user: pd.Series, time: pd.Series, timeout_s: int) -> tuple[pd.Series, pd.Series]:
    """(session_start 0/1, session_id per user). Frame must be in sort order.
    A gap strictly greater than the timeout starts a new session (DEC-N)."""
    gap = time.groupby(user).diff().dt.total_seconds()
    start = (gap.isna() | (gap > timeout_s)).astype("int8")
    return start, start.groupby(user).cumsum()


def last_page_index(is_page: pd.Series, keys: list[pd.Series]) -> pd.Series:
    """For every row, the index of the most recent row with is_page True in the
    same group (the row itself if it is a page). NaN where none exists yet.

    Used for an audio clip's page: shift(1) is not enough, because clips often
    follow other clips (spec §5 var 33)."""
    idx = pd.Series(np.where(is_page, is_page.index, np.nan), index=is_page.index)
    return idx.groupby(keys).ffill()
